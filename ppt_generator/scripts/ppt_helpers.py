"""
PPT Helper Functions
Reusable functions for generating professional PowerPoint presentations.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os


# ── Color Themes ──────────────────────────────────────────────────────────────

DARK_THEME = {
    "bg_primary":     RGBColor(0x1A, 0x1A, 0x2E),
    "bg_secondary":   RGBColor(0x16, 0x21, 0x3E),
    "accent":         RGBColor(0x00, 0xD2, 0xFF),
    "accent_warm":    RGBColor(0xFF, 0x6B, 0x6B),
    "accent_green":   RGBColor(0x4E, 0xCB, 0x71),
    "accent_purple":  RGBColor(0xA8, 0x55, 0xF7),
    "accent_orange":  RGBColor(0xFF, 0x9F, 0x43),
    "text_primary":   RGBColor(0xFF, 0xFF, 0xFF),
    "text_secondary": RGBColor(0xB0, 0xB0, 0xC0),
    "text_muted":     RGBColor(0x80, 0x80, 0x99),
}

LIGHT_THEME = {
    "bg_primary":     RGBColor(0xFF, 0xFF, 0xFF),
    "bg_secondary":   RGBColor(0xF4, 0xF5, 0xF7),
    "accent":         RGBColor(0x2D, 0x5B, 0xE3),
    "accent_warm":    RGBColor(0xE8, 0x4D, 0x4D),
    "accent_green":   RGBColor(0x0D, 0x94, 0x88),
    "accent_purple":  RGBColor(0x7C, 0x3A, 0xED),
    "accent_orange":  RGBColor(0xEA, 0x88, 0x0E),
    "text_primary":   RGBColor(0x1A, 0x1A, 0x2E),
    "text_secondary": RGBColor(0x4A, 0x4A, 0x5A),
    "text_muted":     RGBColor(0x9A, 0x9A, 0xAA),
}


def get_theme(name="dark"):
    """Return theme colors by name."""
    return DARK_THEME if name == "dark" else LIGHT_THEME


# ── Slide Builders ────────────────────────────────────────────────────────────

def _set_bg(slide, color):
    """Set a solid background color on a slide."""
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def _add_accent_bar(slide, left, top, width, height, color):
    """Add a thin decorative accent bar."""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_title_slide(prs, title, subtitle="", theme="dark"):
    """Create a visually striking title slide."""
    t = get_theme(theme)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg(slide, t["bg_primary"])

    # Decorative accent bar
    _add_accent_bar(slide, Inches(0.8), Inches(3.15), Inches(2.0), Inches(0.05), t["accent"])

    # Title
    txBox = slide.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(8.4), Inches(1.6))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(42)
    p.font.bold = True
    p.font.color.rgb = t["text_primary"]

    # Subtitle
    if subtitle:
        txBox2 = slide.shapes.add_textbox(Inches(0.8), Inches(3.4), Inches(8.4), Inches(1.0))
        tf2 = txBox2.text_frame
        tf2.word_wrap = True
        p2 = tf2.paragraphs[0]
        p2.text = subtitle
        p2.font.size = Pt(18)
        p2.font.color.rgb = t["text_secondary"]

    return slide


def add_section_slide(prs, section_title, section_number=None, theme="dark"):
    """Create a section divider slide."""
    t = get_theme(theme)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg(slide, t["bg_primary"])

    if section_number is not None:
        num_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(2.0), Inches(0.8))
        ntf = num_box.text_frame
        np_ = ntf.paragraphs[0]
        np_.text = f"0{section_number}" if section_number < 10 else str(section_number)
        np_.font.size = Pt(48)
        np_.font.bold = True
        np_.font.color.rgb = t["accent"]

    y_offset = Inches(2.8) if section_number else Inches(2.5)
    txBox = slide.shapes.add_textbox(Inches(0.8), y_offset, Inches(8.4), Inches(1.2))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = section_title
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = t["text_primary"]

    _add_accent_bar(slide, Inches(0.8), y_offset + Inches(1.2), Inches(1.5), Inches(0.05), t["accent"])

    return slide


def add_content_slide(prs, title, bullets, theme="dark"):
    """Create a content slide with bullet points."""
    t = get_theme(theme)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg(slide, t["bg_secondary"])

    # Title
    txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8.4), Inches(0.8))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = t["text_primary"]

    _add_accent_bar(slide, Inches(0.8), Inches(1.18), Inches(1.5), Inches(0.04), t["accent"])

    # Bullets
    txBox2 = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(8.4), Inches(4.5))
    tf2 = txBox2.text_frame
    tf2.word_wrap = True
    for i, bullet in enumerate(bullets):
        if i == 0:
            p = tf2.paragraphs[0]
        else:
            p = tf2.add_paragraph()
        p.text = f"•  {bullet}"
        p.font.size = Pt(17)
        p.font.color.rgb = t["text_secondary"]
        p.space_after = Pt(14)

    return slide


def add_two_column_slide(prs, title, left_title, left_items, right_title, right_items, theme="dark"):
    """Create a two-column comparison slide."""
    t = get_theme(theme)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg(slide, t["bg_secondary"])

    # Slide title
    txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8.4), Inches(0.8))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = t["text_primary"]

    _add_accent_bar(slide, Inches(0.8), Inches(1.18), Inches(1.5), Inches(0.04), t["accent"])

    # Left column header
    lh = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(4.0), Inches(0.5))
    lhp = lh.text_frame.paragraphs[0]
    lhp.text = left_title
    lhp.font.size = Pt(18)
    lhp.font.bold = True
    lhp.font.color.rgb = t["accent"]

    # Left column items
    lb = slide.shapes.add_textbox(Inches(0.8), Inches(2.1), Inches(4.0), Inches(3.9))
    lbtf = lb.text_frame
    lbtf.word_wrap = True
    for i, item in enumerate(left_items):
        p = lbtf.paragraphs[0] if i == 0 else lbtf.add_paragraph()
        p.text = f"•  {item}"
        p.font.size = Pt(15)
        p.font.color.rgb = t["text_secondary"]
        p.space_after = Pt(10)

    # Vertical divider
    _add_accent_bar(slide, Inches(5.0), Inches(1.5), Inches(0.02), Inches(4.2), t["text_muted"])

    # Right column header
    rh = slide.shapes.add_textbox(Inches(5.3), Inches(1.5), Inches(4.0), Inches(0.5))
    rhp = rh.text_frame.paragraphs[0]
    rhp.text = right_title
    rhp.font.size = Pt(18)
    rhp.font.bold = True
    rhp.font.color.rgb = t["accent_warm"]

    # Right column items
    rb = slide.shapes.add_textbox(Inches(5.3), Inches(2.1), Inches(4.0), Inches(3.9))
    rbtf = rb.text_frame
    rbtf.word_wrap = True
    for i, item in enumerate(right_items):
        p = rbtf.paragraphs[0] if i == 0 else rbtf.add_paragraph()
        p.text = f"•  {item}"
        p.font.size = Pt(15)
        p.font.color.rgb = t["text_secondary"]
        p.space_after = Pt(10)

    return slide


def add_table_slide(prs, title, headers, rows, theme="dark"):
    """Create a slide with a styled table."""
    t = get_theme(theme)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg(slide, t["bg_secondary"])

    # Title
    txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8.4), Inches(0.8))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = t["text_primary"]

    _add_accent_bar(slide, Inches(0.8), Inches(1.18), Inches(1.5), Inches(0.04), t["accent"])

    # Table
    num_rows = len(rows) + 1  # +1 for header
    num_cols = len(headers)
    table_width = Inches(8.4)
    table_height = Inches(0.4 * num_rows)
    table_shape = slide.shapes.add_table(
        num_rows, num_cols, Inches(0.8), Inches(1.5), table_width, table_height
    )
    table = table_shape.table

    # Header row
    for j, header in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = header
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.size = Pt(14)
            paragraph.font.bold = True
            paragraph.font.color.rgb = t["text_primary"]
        cell.fill.solid()
        cell.fill.fore_color.rgb = t["accent"]

    # Data rows
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.cell(i + 1, j)
            cell.text = str(value)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(13)
                paragraph.font.color.rgb = t["text_secondary"]
            cell.fill.solid()
            cell.fill.fore_color.rgb = t["bg_primary"] if i % 2 == 0 else t["bg_secondary"]

    return slide


def add_stats_slide(prs, title, stats, theme="dark"):
    """
    Create a slide showcasing key statistics/metrics.
    stats: list of dicts with 'value' and 'label' keys.
    """
    t = get_theme(theme)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg(slide, t["bg_secondary"])

    # Title
    txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8.4), Inches(0.8))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = t["text_primary"]

    _add_accent_bar(slide, Inches(0.8), Inches(1.18), Inches(1.5), Inches(0.04), t["accent"])

    # Stats cards
    num_stats = len(stats)
    card_width = 8.4 / num_stats
    accent_colors = [t["accent"], t["accent_warm"], t["accent_green"], t["accent_purple"], t["accent_orange"]]

    for i, stat in enumerate(stats):
        x = Inches(0.8 + i * card_width)

        # Value
        vbox = slide.shapes.add_textbox(x, Inches(2.2), Inches(card_width - 0.2), Inches(1.0))
        vtf = vbox.text_frame
        vp = vtf.paragraphs[0]
        vp.text = str(stat["value"])
        vp.font.size = Pt(40)
        vp.font.bold = True
        vp.font.color.rgb = accent_colors[i % len(accent_colors)]
        vp.alignment = PP_ALIGN.CENTER

        # Label
        lbox = slide.shapes.add_textbox(x, Inches(3.2), Inches(card_width - 0.2), Inches(0.6))
        ltf = lbox.text_frame
        ltf.word_wrap = True
        lp = ltf.paragraphs[0]
        lp.text = stat["label"]
        lp.font.size = Pt(14)
        lp.font.color.rgb = t["text_muted"]
        lp.alignment = PP_ALIGN.CENTER

    return slide


def add_thank_you_slide(prs, message="Thank You", contact="", theme="dark"):
    """Create a closing/thank-you slide."""
    t = get_theme(theme)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg(slide, t["bg_primary"])

    # Message
    txBox = slide.shapes.add_textbox(Inches(1), Inches(2.2), Inches(8), Inches(1.5))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = message
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = t["accent"]
    p.alignment = PP_ALIGN.CENTER

    _add_accent_bar(slide, Inches(4.0), Inches(3.7), Inches(2.0), Inches(0.04), t["accent"])

    # Contact info
    if contact:
        txBox2 = slide.shapes.add_textbox(Inches(1), Inches(4.0), Inches(8), Inches(1.0))
        tf2 = txBox2.text_frame
        tf2.word_wrap = True
        p2 = tf2.paragraphs[0]
        p2.text = contact
        p2.font.size = Pt(16)
        p2.font.color.rgb = t["text_muted"]
        p2.alignment = PP_ALIGN.CENTER

    return slide


def create_presentation(widescreen=True):
    """Create a new Presentation object with standard dimensions."""
    prs = Presentation()
    if widescreen:
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
    return prs


def save_presentation(prs, filename, directory=None):
    """Save the presentation and print the output path."""
    if directory is None:
        directory = os.path.join(os.path.expanduser("~"), "Desktop")
    os.makedirs(directory, exist_ok=True)
    if not filename.endswith(".pptx"):
        filename += ".pptx"
    output_path = os.path.join(directory, filename)
    prs.save(output_path)
    print(f"✅ Presentation saved to: {output_path}")
    return output_path
